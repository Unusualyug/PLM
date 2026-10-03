import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from data_part1 import SESSIONS_1_18
from data_part2 import SESSIONS_19_36
from data_part3 import SESSIONS_37_53

SESSIONS = SESSIONS_1_18 + SESSIONS_19_36 + SESSIONS_37_53

UNIV = "Gyanmanjari Innovative University"
COLLEGE = "Gyanmanjari Degree Engineering College"
DEPT = "Department of Computer Science and Engineering"
SEM = "Semester - 3"
SUBJECT = "Database Management System: Organizing, Querying, and Managing Data - BET2CS13305"
FACULTY = "Faculty Members: Dr. Mansi C. Shanishwara (MCS), Dipendrasinh P. Zala (DPZ), Yug G. Lakhani (YGL)"
HRS = "Hrs./Week: 6 Hours"
TITLE = "Revised Micro Session Plan with Detailed SOP and Session-Wise Question Examples (PREMIUM / PLM Class)"


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), "true")
    trPr.append(tblHeader)


doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Inches(16.54), Inches(11.69)  # A3 landscape for readability
sec.left_margin = sec.right_margin = Inches(0.4)
sec.top_margin = sec.bottom_margin = Inches(0.4)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(9)

def head(text, size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    return p

# ---------------- Cover / header block ----------------
head(UNIV, 16)
head(COLLEGE, 13)
head(DEPT, 12)
head(SEM, 11)
head(SUBJECT, 12, space=4)
head(TITLE, 12, space=6)

info = doc.add_table(rows=1, cols=2)
info.style = 'Table Grid'
info.rows[0].cells[0].text = FACULTY
info.rows[0].cells[1].text = HRS
for c in info.rows[0].cells:
    c.paragraphs[0].runs[0].font.size = Pt(10)
    c.paragraphs[0].runs[0].bold = True

doc.add_paragraph()
head("Revision Note (as per the directions of Provost Sir, communicated on 03-10-2026)", 11, align=WD_ALIGN_PARAGRAPH.LEFT, space=4)
notes = [
 "This session plan has been revised specifically for PREMIUM / PLM pedagogy. It is not a conventional lecture plan.",
 "For every session, three additional components have been added: (1) Session Execution Plan, (2) Detailed SOP for the session, and (3) Real-life / unsolved / open-book examination-oriented questions.",
 "Each session clearly demonstrates what students will do, how the faculty will execute the session, and what students are expected to achieve (stated as the deliverable at the end of every SOP).",
 "All questions are topic-specific and practically relevant. Generic, recall-only questions have been avoided and replaced with real-world scenarios, industry problems, case studies, practical situations, open-book tasks and application-oriented problems.",
 "Wherever a real-time, industry-based or senior-level example is used in a session, a corresponding question or problem based on that example has been provided in the same session.",
 "The question standard reflects the type and level of questions expected in the PLM Sessional Examinations - they require application, analysis, problem-solving and practical thinking.",
 "Faculty members handling this subject (MCS, DPZ, YGL) will follow this common plan, common SOP structure and common question standard to maintain complete consistency across all divisions.",
]
for n in notes:
    p = doc.add_paragraph(n, style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    for r in p.runs:
        r.font.size = Pt(9)

doc.add_page_break()

# ---------------- Main table ----------------
HEADERS = ["Session No. / Module", "Topics & Module Description (Time / Teacher Aid)",
           "Session Execution Plan (How the session will be conducted)",
           "Detailed SOP for the Session (Faculty & Student Activity Sequence + Expected Outcome)",
           "Real-Life / Unsolved / Open-Book Examination-Oriented Questions",
           "Planned Date", "Actual Date"]
WIDTHS = [0.85, 3.2, 3.3, 4.3, 3.4, 0.7, 0.7]

table = doc.add_table(rows=1, cols=len(HEADERS))
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table.rows[0]
set_repeat_header(hdr)
for i, h in enumerate(HEADERS):
    cell = hdr.cells[i]
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.bold = True
    r.font.size = Pt(9)
    shade(cell, "1F4E79")
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

current_module = None
MODULE_TITLES = {
 "1": "Module-1 : Introduction to Database Management Concept and SQL Basics",
 "2": "Module-2 : Database System Architecture and Ensuring Data Integrity",
 "3": "Module-3 : Advanced SQL and Normalization",
 "4": "Module-4 : PL/SQL and Database Security",
 "5": "Module-5 : Advanced Concepts of SQL & Access Control",
}

def add_module_row(title):
    row = table.add_row()
    cells = row.cells
    merged = cells[0]
    for c in cells[1:]:
        merged = merged.merge(c)
    merged.text = ''
    p = merged.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10)
    shade(merged, "DDEBF7")

def put(cell, blocks, size=8):
    """blocks: list of (text, bold, is_bullet)"""
    cell.text = ''
    first = True
    for text, bold, bullet in blocks:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(2)
        if bullet:
            p.paragraph_format.left_indent = Inches(0.08)
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)

for s in SESSIONS:
    mod = str(s["module"])
    if mod in MODULE_TITLES and mod != current_module:
        add_module_row(MODULE_TITLES[mod])
        current_module = mod

    row = table.add_row()
    c = row.cells

    put(c[0], [(f"Session {s['no']}", True, False), (f"Module {s['module']}", False, False)], 9)

    blocks = [(f"Topic: {s['title']}", True, False)]
    for d, t, aid in s["breakdown"]:
        blocks.append((f"\u2022 {d}  |  {t}  |  {aid}", False, True))
    if s.get("homework"):
        blocks.append(("Homework:", True, False))
        for i, h in enumerate(s["homework"], 1):
            blocks.append((f"{i}. {h}", False, True))
    put(c[1], blocks)

    put(c[2], [(s["exec"], False, False)])

    sop_blocks = []
    for step in s["sop"]:
        bold = step.lower().startswith("deliverable")
        sop_blocks.append((step, bold, True))
    put(c[3], sop_blocks)

    q_blocks = []
    for i, q in enumerate(s["questions"], 1):
        q_blocks.append((f"Q{i}. {q}", False, True))
    put(c[4], q_blocks)

    c[5].text = ""
    c[6].text = ""

for i, w in enumerate(WIDTHS):
    for row in table.rows:
        row.cells[i].width = Inches(w)

doc.add_paragraph()
sign = doc.add_table(rows=1, cols=2)
sign.rows[0].cells[0].text = "\n\nSignature of Subject Faculty\n(MCS / DPZ / YGL)"
sign.rows[0].cells[1].text = "\n\nSignature of Head of Department"
sign.rows[0].cells[1].paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.RIGHT

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Revised_Micro_Session_Plan_DBMS_PLM.docx")
doc.save(out)
print("saved:", out, os.path.getsize(out))
print("sessions:", len(SESSIONS))
