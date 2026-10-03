import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from data_part1 import SESSIONS_1_18
from data_part2 import SESSIONS_19_36
from data_part3 import SESSIONS_37_53

S = SESSIONS_1_18 + SESSIONS_19_36 + SESSIONS_37_53

NAVY   = RGBColor(0x1F, 0x4E, 0x79)
MODFILL= "D9E1F2"
HDRFILL= "D9D9D9"
NEW_1  = "FFF2CC"
NEW_2  = "E2EFDA"
NEW_3  = "FCE4EC"
TOPIC  = "F2F2F2"

HEADERS = ["Session no.", "Module", "Description", "Time", "Teacher aid",
           "Planned date", "Actual date", "Remarks"]
WIDTHS  = [0.75, 0.6, 8.1, 0.75, 1.9, 0.95, 0.95, 1.1]


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def set_repeat(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), "true"); trPr.append(el)


def write(cell, text, size=8.5, bold=False, align=None, color=None, fill=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    if align: p.alignment = align
    first = True
    for line in str(text).split("\n"):
        if not first:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            if align: p.alignment = align
        first = False
        r = p.add_run(line)
        r.font.size = Pt(size); r.bold = bold
        if color: r.font.color.rgb = color
    if fill: shade(cell, fill)


doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Inches(16.54), Inches(11.69)   # A3 landscape
sec.left_margin = sec.right_margin = Inches(0.35)
sec.top_margin = sec.bottom_margin = Inches(0.35)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(9)


def banner(text, size, bold=True, color=None, space=0, italic=False):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space)
    r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = color


banner("Gyanmanjari Innovative University", 16, color=NAVY)
banner("Gyanmanjari Degree Engineering College", 13)
banner("Department of Computer Science and Engineering", 12)
banner("Semester-3", 11)
banner("Database Management System: Organizing, Querying, and Managing Data \u2013 BET2CS13305", 12.5, color=NAVY)
banner("Micro Session Planner", 13, space=2)
banner("Name of Faculty: Dr. Mansi C. Shanishwara (MCS) / Dipendrasinh P. Zala (DPZ) / Yug G. Lakhani (YGL)"
       "                    Hrs. Week: 6 hours", 10.5, space=2)
banner("REVISED AS PER DIRECTIONS OF PROVOST SIR \u2013 WITH SESSION EXECUTION PLAN, DETAILED SOP AND "
       "SESSION-WISE QUESTION EXAMPLES (PREMIUM / PLM CLASS)", 11,
       color=RGBColor(0xC0, 0x00, 0x00), space=4)

# ---- colour legend ----
leg = doc.add_table(rows=1, cols=4)
leg.style = 'Table Grid'
leg.alignment = WD_TABLE_ALIGNMENT.CENTER
items = [("NEWLY ADDED \u2013 1. Session Execution Plan", NEW_1),
         ("NEWLY ADDED \u2013 2. Detailed SOP (with Deliverable)", NEW_2),
         ("NEWLY ADDED \u2013 3. Real-Life / Open-Book Questions", NEW_3),
         ("Original content (unchanged)", "FFFFFF")]
for i, (t, f) in enumerate(items):
    write(leg.rows[0].cells[i], t, size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, fill=f)
    leg.rows[0].cells[i].width = Inches(3.8)
doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ---- main table ----
table = doc.add_table(rows=1, cols=len(HEADERS))
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
hdr = table.rows[0]
set_repeat(hdr)
for i, h in enumerate(HEADERS):
    write(hdr.cells[i], h, size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, fill=HDRFILL)

MODULE_TITLES = {
 "1": "Module-1 : Introduction to Database Management Concept and SQL Basics",
 "2": "Module-2 : Database System Architecture and Ensuring Data Integrity",
 "3": "Module-3 : Advanced SQL and Normalization",
 "4": "Module-4 : PL/SQL and Database Security",
 "5": "Module-5 : Advanced Concepts of SQL & Access Control",
}

def add_module(title):
    row = table.add_row()
    m = row.cells[0]
    for c in row.cells[1:]:
        m = m.merge(c)
    write(m, title, size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY, fill=MODFILL)


def activity_row(desc, time_, aid):
    row = table.add_row()
    c = row.cells
    write(c[2], desc)
    write(c[3], time_, align=WD_ALIGN_PARAGRAPH.CENTER)
    write(c[4], aid, align=WD_ALIGN_PARAGRAPH.CENTER)
    return row


def merged_row(text, fill=None, bold=False, size=8.5, color=None):
    row = table.add_row()
    c = row.cells
    m = c[2].merge(c[3]).merge(c[4])
    write(m, text, size=size, bold=bold, color=color, fill=fill)
    return row


current = None
for s in S:
    mod = str(s["module"])
    if mod in MODULE_TITLES and mod != current:
        add_module(MODULE_TITLES[mod]); current = mod

    rows = []
    rows.append(merged_row(f"Topics: {s['title']}", fill=TOPIC, bold=True, size=9.5))
    for d, t, aid in s["breakdown"]:
        rows.append(activity_row(d, t, aid))
    if s.get("homework"):
        rows.append(merged_row("Homework:\n" + "\n".join(f"{i}. {h}" for i, h in enumerate(s["homework"], 1)),
                               fill=TOPIC))

    rows.append(merged_row("1. SESSION EXECUTION PLAN  (how the session will be conducted / executed)",
                           fill=NEW_1, bold=True, size=9.5, color=RGBColor(0x7F, 0x60, 0x00)))
    rows.append(merged_row(s["exec"], fill=NEW_1))

    rows.append(merged_row("2. DETAILED SOP FOR THE SESSION  (sequence of faculty & student activities + expected outcome / deliverable)",
                           fill=NEW_2, bold=True, size=9.5, color=RGBColor(0x37, 0x56, 0x23)))
    for step in s["sop"]:
        rows.append(merged_row(step, fill=NEW_2, bold=step.lower().startswith("deliverable")))

    rows.append(merged_row("3. REAL-LIFE / UNSOLVED / CASE-BASED / OPEN-BOOK EXAMINATION-ORIENTED QUESTIONS",
                           fill=NEW_3, bold=True, size=9.5, color=RGBColor(0x7B, 0x2E, 0x43)))
    for i, q in enumerate(s["questions"], 1):
        rows.append(merged_row(f"Q{i}. {q}", fill=NEW_3))

    # vertical merge of Session no., Module, Planned/Actual date, Remarks across the block
    for col in (0, 1, 5, 6, 7):
        target = rows[0].cells[col]
        for rw in rows[1:]:
            target = target.merge(rw.cells[col])
    write(rows[0].cells[0], str(s["no"]), size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write(rows[0].cells[1], str(s["module"]), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

for row in table.rows:
    for i, w in enumerate(WIDTHS):
        try:
            row.cells[i].width = Inches(w)
        except IndexError:
            pass

doc.add_paragraph()
sign = doc.add_table(rows=1, cols=2)
write(sign.rows[0].cells[0], "\n\nSignature of Subject Faculty   (MCS / DPZ / YGL)", size=11, bold=True)
write(sign.rows[0].cells[1], "\n\nSignature of Head of Department", size=11, bold=True,
      align=WD_ALIGN_PARAGRAPH.RIGHT)

out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "Revised_Micro_Session_Plan_DBMS_PLM_ORIGINAL_FORMAT.docx")
doc.save(out)
print("saved", out, os.path.getsize(out))
