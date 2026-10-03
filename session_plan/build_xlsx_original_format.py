import sys, os, math
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from data_part1 import SESSIONS_1_18
from data_part2 import SESSIONS_19_36
from data_part3 import SESSIONS_37_53

S = SESSIONS_1_18 + SESSIONS_19_36 + SESSIONS_37_53

# ---------------- styling constants ----------------
NAVY   = "1F4E79"
MODFILL= "D9E1F2"   # module band (original style)
HDRFILL= "D9D9D9"   # grey header like original printed sheet
NEW_1  = "FFF2CC"   # NEW: Session Execution Plan  (light amber)
NEW_2  = "E2EFDA"   # NEW: Detailed SOP            (light green)
NEW_3  = "FCE4EC"   # NEW: Questions               (light pink)
TOPIC  = "F2F2F2"
WHITE  = "FFFFFF"

thin  = Side(style="thin", color="A6A6A6")
BORDER= Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP  = Alignment(wrap_text=True, vertical="top")
CTR   = Alignment(wrap_text=True, vertical="center", horizontal="center")

# original column layout
HEADERS = ["Session no.", "Module", "Description", "Time", "Teacher aid",
           "Planned date", "Actual date", "Remarks"]
WIDTHS  = [9.5, 8, 96, 9.5, 24, 12, 12, 14]
DESC_CHARS = int(WIDTHS[2] * 1.08)          # chars per line in Description column
MERGE_CHARS = int((WIDTHS[2] + WIDTHS[3] + WIDTHS[4]) * 1.06)

wb = Workbook()
ws = wb.active
ws.title = "Micro Session Planner"
ws.sheet_view.showGridLines = False

r = 1
def banner(text, size, bold=True, color="000000", height=20, fill=None, italic=False):
    global r
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HEADERS))
    c = ws.cell(row=r, column=1, value=text)
    c.font = Font(bold=bold, size=size, color=color, italic=italic)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if fill:
        for col in range(1, len(HEADERS) + 1):
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=fill)
    ws.row_dimensions[r].height = height
    r += 1

# ---------------- original title block ----------------
banner("Gyanmanjari Innovative University", 15, color=NAVY, height=22)
banner("Gyanmanjari Degree Engineering College", 12.5, height=18)
banner("Department of Computer Science and Engineering", 11.5, height=17)
banner("Semester-3", 11, height=16)
banner("Database Management System: Organizing, Querying, and Managing Data \u2013 BET2CS13305", 12, color=NAVY, height=18)
banner("Micro Session Planner", 13, height=19)
banner("Name of Faculty: Dr. Mansi C. Shanishwara (MCS) / Dipendrasinh P. Zala (DPZ) / Yug G. Lakhani (YGL)"
       "                    Hrs. Week: 6 hours", 10.5, height=17)
banner("REVISED AS PER DIRECTIONS OF PROVOST SIR \u2013 WITH SESSION EXECUTION PLAN, DETAILED SOP AND SESSION-WISE QUESTION EXAMPLES (PREMIUM / PLM CLASS)",
       11, color="C00000", height=20)

# ---------------- colour legend ----------------
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
c = ws.cell(row=r, column=1, value="COLOUR LEGEND \u2192")
c.font = Font(bold=True, size=10); c.alignment = CTR; c.border = BORDER
legend = [("NEWLY ADDED \u2013 1. Session Execution Plan", NEW_1),
          ("NEWLY ADDED \u2013 2. Detailed SOP (with Deliverable)", NEW_2),
          ("NEWLY ADDED \u2013 3. Real-Life / Open-Book Questions", NEW_3),
          ("Original content (unchanged)", WHITE)]
col = 3
for text, fill in legend:
    if col > len(HEADERS):
        break
    end = min(col + 1, len(HEADERS))
    ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=end)
    cc = ws.cell(row=r, column=col, value=text)
    cc.font = Font(bold=True, size=9)
    cc.alignment = CTR
    for k in range(col, end + 1):
        ws.cell(row=r, column=k).fill = PatternFill("solid", fgColor=fill)
        ws.cell(row=r, column=k).border = BORDER
    col = end + 1
ws.row_dimensions[r].height = 26
r += 2

# ---------------- column header row (as in original) ----------------
HDR_ROW = r
for i, h in enumerate(HEADERS, 1):
    c = ws.cell(row=HDR_ROW, column=i, value=h)
    c.font = Font(bold=True, size=10.5)
    c.fill = PatternFill("solid", fgColor=HDRFILL)
    c.alignment = CTR
    c.border = BORDER
ws.row_dimensions[HDR_ROW].height = 24
r += 1

MODULE_TITLES = {
 "1": "Module-1 : Introduction to Database Management Concept and SQL Basics",
 "2": "Module-2 : Database System Architecture and Ensuring Data Integrity",
 "3": "Module-3 : Advanced SQL and Normalization",
 "4": "Module-4 : PL/SQL and Database Security",
 "5": "Module-5 : Advanced Concepts of SQL & Access Control",
}

def h_for(text, chars, base=13.2, pad=2):
    lines = 0
    for para in str(text).split("\n"):
        lines += max(1, math.ceil(len(para) / chars))
    return lines * base + pad

def activity_row(desc, time_, aid, fill=None, bold=False, size=9.5):
    """normal 3-column row: Description | Time | Teacher aid"""
    global r
    vals = [None, None, desc, time_, aid, None, None, None]
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.border = BORDER
        c.alignment = WRAP if i == 3 else CTR
        c.font = Font(size=size, bold=bold)
        if fill and i in (3, 4, 5):
            c.fill = PatternFill("solid", fgColor=fill)
    ws.row_dimensions[r].height = h_for(desc, DESC_CHARS)
    r += 1
    return r - 1

def merged_row(text, fill=None, bold=False, size=9.5, italic=False, color="000000"):
    """row where Description+Time+Teacher aid are merged (used for Topics / Homework /
       and for the newly added Execution Plan, SOP and Question content)"""
    global r
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    c = ws.cell(row=r, column=3, value=text)
    c.alignment = WRAP
    c.font = Font(size=size, bold=bold, italic=italic, color=color)
    for i in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=r, column=i)
        cell.border = BORDER
        if fill and i in (3, 4, 5):
            cell.fill = PatternFill("solid", fgColor=fill)
    ws.row_dimensions[r].height = h_for(text, MERGE_CHARS)
    r += 1
    return r - 1

current_mod = None
for s in S:
    mod = str(s["module"])
    if mod in MODULE_TITLES and mod != current_mod:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HEADERS))
        c = ws.cell(row=r, column=1, value=MODULE_TITLES[mod])
        c.font = Font(bold=True, size=11, color=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        for col in range(1, len(HEADERS) + 1):
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=MODFILL)
            ws.cell(row=r, column=col).border = BORDER
        ws.row_dimensions[r].height = 20
        current_mod = mod
        r += 1

    block_start = r

    # --- ORIGINAL CONTENT ---
    merged_row(f"Topics: {s['title']}", fill=TOPIC, bold=True, size=10)
    for d, t, aid in s["breakdown"]:
        activity_row(d, t, aid)
    if s.get("homework"):
        merged_row("Homework:\n" + "\n".join(f"{i}. {h}" for i, h in enumerate(s["homework"], 1)),
                   fill=TOPIC)

    # --- NEWLY ADDED CONTENT (highlighted) ---
    merged_row("1. SESSION EXECUTION PLAN  (how the session will be conducted / executed)",
               fill=NEW_1, bold=True, size=10, color="7F6000")
    merged_row(s["exec"], fill=NEW_1)

    merged_row("2. DETAILED SOP FOR THE SESSION  (sequence of faculty & student activities + expected outcome / deliverable)",
               fill=NEW_2, bold=True, size=10, color="375623")
    for step in s["sop"]:
        merged_row(step, fill=NEW_2, bold=step.lower().startswith("deliverable"))

    merged_row("3. REAL-LIFE / UNSOLVED / CASE-BASED / OPEN-BOOK EXAMINATION-ORIENTED QUESTIONS",
               fill=NEW_3, bold=True, size=10, color="7B2E43")
    for i, q in enumerate(s["questions"], 1):
        merged_row(f"Q{i}. {q}", fill=NEW_3)

    block_end = r - 1

    # --- merge Session no. / Module / dates / remarks across the whole block ---
    ws.merge_cells(start_row=block_start, start_column=1, end_row=block_end, end_column=1)
    ws.merge_cells(start_row=block_start, start_column=2, end_row=block_end, end_column=2)
    for col in (6, 7, 8):
        ws.merge_cells(start_row=block_start, start_column=col, end_row=block_end, end_column=col)
    c = ws.cell(row=block_start, column=1, value=s["no"])
    c.font = Font(bold=True, size=12); c.alignment = CTR
    c2 = ws.cell(row=block_start, column=2, value=s["module"])
    c2.font = Font(size=11); c2.alignment = CTR
    for col in (6, 7, 8):
        ws.cell(row=block_start, column=col).alignment = CTR

    # thicker separation line at the end of each session block
    med = Side(style="medium", color="808080")
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=block_end, column=col)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=med)

# ---------------- signature block (as in original) ----------------
r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
c = ws.cell(row=r, column=1, value="Signature of Subject Faculty   (MCS / DPZ / YGL)")
c.font = Font(bold=True, size=11)
ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
c = ws.cell(row=r, column=5, value="Signature of Head of Department")
c.font = Font(bold=True, size=11); c.alignment = Alignment(horizontal="right")
ws.row_dimensions[r].height = 45

for i, w in enumerate(WIDTHS, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = f"C{HDR_ROW + 1}"
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = 8          # A3
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = f"{HDR_ROW}:{HDR_ROW}"
ws.print_options.horizontalCentered = True

out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "Revised_Micro_Session_Plan_DBMS_PLM_ORIGINAL_FORMAT.xlsx")
wb.save(out)
print("saved", out, os.path.getsize(out), "| last row:", r)
